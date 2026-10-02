"""책자 PPT 생성 — 양식(첨부1) 위에 7장. 실행: python3 build_book.py
스크린샷 교체: SHOTS 의 경로만 바꾸고 다시 실행.
"""
import os
import sys
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from pptlib import *  # noqa

ROOT = os.path.abspath(os.path.join(HERE, "../../../.."))
DOCS = os.path.join(ROOT, "docs")
TEMPLATE = os.path.join(DOCS, "deliverables/templates/첨부1_책자양식.pptx")
OUT = os.path.join(DOCS, "deliverables/book/MARU_책자.pptx")
ASSETS = os.path.join(HERE, "assets")
os.makedirs(ASSETS, exist_ok=True)

CH = os.path.join(DOCS, "features/character/screenshots")
MS = os.path.join(DOCS, "features/mission/screenshots")
FINAL = os.path.join(DOCS, "deliverables/screenshots")

# 임시 스크린샷 → 최종 키트(docs/deliverables/screenshots) 나오면 교체
SHOTS = {
    "lesson_hint": f"{CH}/apply_lesson_C1.png",
    "lesson_done": f"{CH}/apply_lesson_C3.png",
    "vocab_review": f"{CH}/apply_vocab_C3_review_done.png",
    "vocab_match": f"{CH}/apply_vocab_C2.png",
    "mission_magic": f"{MS}/r3_fe2/12_loading_magic.png",
    "mission_red": f"{MS}/MSN-1.7.3_red_card.png",
    "mission_report": f"{MS}/r3_fe2/8_report_p1_not_cleared.png",
    "lab_grammar": f"{CH}/apply_lab_C3.png",
}
RAW = os.path.join(DOCS, "features/character/raw")

HEADER = "2026년 졸업작품 · MARU (마루)"
FOOTER = "팀원 하도윤 · 지도교수 한경수 교수"
GW = 10.0


def prep_assets():
    ph, ch = {}, {}
    for k, p in SHOTS.items():
        ph[k] = phone_frame(p, f"{ASSETS}/ph_{k}.png", width=640)
    for f in os.listdir(RAW):
        if f.endswith(".png"):
            ch[f[:-4]] = trim_alpha(f"{RAW}/{f}", f"{ASSETS}/c_{f}", max_side=700)
    return ph, ch


def main():
    ph, ch = prep_assets()
    prs = Presentation(TEMPLATE)
    C = Canvas(prs, GW)
    GH = prs.slide_height / C.unit  # ≈ 7.48

    # 구조: 예시(1번) 삭제 → 6장 + 1장 추가(제목 및 내용) → 7장
    layout = prs.slides[1].slide_layout
    prs.slides.add_slide(layout)   # 파트 이름 충돌 방지: 추가 먼저, 삭제 나중
    move_slide(prs, 7, 6)          # 새 장을 '결론' 앞(캐릭터)으로
    delete_slide(prs, 0)
    slides = list(prs.slides)

    titles = ["필요성 및 배경", "시스템 구성도", "주요 기술 ① Morpheme 레슨",
              "주요 기술 ② Adaptive 단어 (FSRS-4.5)", "주요 기술 ③ 미션 대화 · 언어 실험실",
              "캐릭터 · 사용자 경험", "결론 및 기대효과"]

    for i, s in enumerate(slides):
        title = body = None
        for shp in list(s.placeholders):
            t = shp.placeholder_format.type
            if t == 1:
                title = shp
            else:
                remove_shape(shp)
        # 제목
        title.left, title.top, title.width, title.height = C.e(1.28), C.e(0.42), C.e(8.2), C.e(0.72)
        tf = title.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = titles[i]
        C.style_run(r, 30, NAVY, True)
        from pptx.enum.text import MSO_ANCHOR
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        # 번호 배지
        C.label(s, 0.5, 0.48, 0.6, 0.6, f"{i + 1:02d}", fill=PURPLE, size=17)
        # 머리말 / 바닥글
        C.text(s, 0.5, 0.12, 6, 0.25, HEADER, size=10, color=GRAY, name="머리말")
        C.text(s, 6.5, 0.12, 3.0, 0.25, f"{i + 1} / 7", size=10, color=GRAY, align="r")
        C.text(s, 0.5, GH - 0.38, 9.0, 0.25, FOOTER, size=10, color=GRAY, align="c", name="바닥글")

    page1(C, slides[0], ch)
    page2(C, slides[1], ch)
    page3(C, slides[2], ph, ch)
    page4(C, slides[3], ph, ch)
    page5(C, slides[4], ph, ch)
    page6(C, slides[5], ph, ch)
    page7(C, slides[6], ch)

    prs.core_properties.title = "MARU (마루) — 졸업작품 책자"
    prs.core_properties.author = "하도윤"
    prs.part._element.set("embedTrueTypeFonts", "1")  # PowerPoint 재저장 시 글꼴 포함
    prs.save(OUT)
    print("saved", OUT)


# ---------------------------------------------------------------- 1
def page1(C, s, ch):
    # 왼쪽: 교착어
    C.box(s, 0.5, 1.4, 4.45, 3.55, fill=LILAC)
    C.text(s, 0.75, 1.58, 4.0, 0.4, "한국어는 ‘붙여서’ 만드는 언어", size=18, bold=True, color=NAVY)
    chips = [("저", 0), ("는", 1), ("커피", 0), ("를", 1), ("좋아해요", 0)]
    widths = {"저": 0.42, "는": 0.42, "커피": 0.62, "를": 0.42, "좋아해요": 1.12}
    x = 0.78
    for j, (t, g) in enumerate(chips):
        w = widths[t]
        C.label(s, x, 2.22, w, 0.52, t, fill=(PURPLE if g else WHITE), color=(WHITE if g else NAVY),
                size=17, line=(None if g else LILAC2), radius=0.1)
        x += w + (0.06 if t in ("저", "커피") else 0.22)
    C.text(s, 1.12, 2.8, 0.9, 0.3, "주제 조사", size=10, color=PURPLE, bold=True, align="c")
    C.text(s, 2.18, 2.8, 0.9, 0.3, "목적 조사", size=10, color=PURPLE, bold=True, align="c")
    C.text(s, 2.85, 2.8, 1.75, 0.3, "해요체(존댓말) 동사", size=10, color=GRAY, align="c")
    C.text(s, 0.78, 3.15, 4.0, 0.3, "“I like coffee.”", size=12, color=GRAY)
    C.text(s, 0.78, 3.5, 4.0, 1.4, [
        "단어 뒤에 조사·어미가 붙어 문장 속 역할이 정해지는 교착어.",
        "영어처럼 어순만 외워서는 구조가 보이지 않는다.",
        [("미국 외교연수원(FSI) 기준, 영어 화자에게 ", {}), ("최고 난이도(Category IV)", {"bold": True, "color": PURPLE}), (" 언어.", {})],
    ], size=12.5, color=INK, line=1.3)

    # 오른쪽: 기존 앱 한계
    C.box(s, 5.15, 1.4, 4.35, 3.55, fill=WHITE, line=LILAC2, line_w=1.5)
    C.text(s, 5.4, 1.58, 3.9, 0.4, "기존 학습 앱의 한계", size=18, bold=True, color=NAVY)
    items = [
        (YELLOW, "암기 · 객관식 중심", "단어 카드와 보기 고르기 위주라 실제로 문장을 만들어 보는 경험이 적다."),
        (MINT, "구조를 체험할 기회 부족", "조사·어미를 직접 붙이고 떼어 보며 원리를 이해하는 과정이 없다."),
        (PINK, "AI 대화는 목적·끝·피드백이 없음", "무엇을 말할지, 언제 끝낼지, 무엇을 고칠지 학습자가 알 수 없다."),
    ]
    y = 2.15
    for col, h, d in items:
        C.box(s, 5.4, y + 0.05, 0.36, 0.36, fill=col, shape=MSO_SHAPE.OVAL)
        C.text(s, 5.92, y, 3.4, 0.35, h, size=14, bold=True, color=NAVY)
        C.text(s, 5.92, y + 0.36, 3.4, 0.55, d, size=11.5, color=GRAY, line=1.25)
        y += 0.92

    # 아래: MARU 의 답
    C.box(s, 0.5, 5.15, 9.0, 1.75, fill=PURPLE)
    C.image(s, ch["rabbit_happy"], 0.7, 5.2, h=1.65)
    C.image(s, ch["turtle_thinking"], 8.0, 5.25, h=1.6)
    C.label(s, 2.35, 5.38, 1.4, 0.38, "MARU의 답", fill=WHITE, color=PURPLE, size=12)
    C.text(s, 2.35, 5.85, 5.9, 0.5, [[("토끼처럼 ", {}), ("빠르게 체험", {"color": YELLOW}),
                                    ("하고, 거북이처럼 ", {}), ("천천히 분석", {"color": MINT}), ("한다", {})]],
           size=16, bold=True, color=WHITE)
    C.text(s, 2.35, 6.38, 5.6, 0.4, "한국어의 구조를 직접 체험하는 학습 앱 — 레슨 · 단어 · 미션 · 실험실",
           size=12.5, color=WHITE)


# ---------------------------------------------------------------- 2
def page2(C, s, ch):
    # 앱
    C.box(s, 0.5, 1.45, 2.75, 3.75, fill=LILAC, line=PURPLE, line_w=1.5)
    C.text(s, 0.7, 1.58, 2.4, 0.35, "Flutter 앱", size=16, bold=True, color=NAVY)
    C.text(s, 0.7, 1.93, 2.4, 0.3, "iOS · Android 단일 코드", size=10.5, color=GRAY)
    feats = [("M", "형태소 레슨"), ("A", "적응형 단어 복습"), ("R", "롤플레이 미션"), ("U", "언어 실험실")]
    y = 2.35
    for k, t in feats:
        C.box(s, 0.7, y, 2.35, 0.42, fill=WHITE, radius=0.1)
        C.label(s, 0.78, y + 0.06, 0.3, 0.3, k, fill=PURPLE, size=11, radius=0.06)
        C.text(s, 1.18, y, 1.8, 0.42, t, size=12, color=NAVY, bold=True, anchor="m")
        y += 0.5
    C.box(s, 0.7, y, 2.35, 0.42, fill=PINK, radius=0.1)
    C.text(s, 0.7, y, 2.35, 0.42, "토끼 · 거북이 캐릭터", size=12, color=WHITE, bold=True, anchor="m", align="c")
    C.text(s, 0.7, y + 0.5, 2.4, 0.3, "Riverpod · Dio · just_audio", size=10, color=GRAY)

    # 앱 ↔ 서버
    C.arrow(s, 3.3, 3.15, 3.95, 3.15)
    C.arrow(s, 3.95, 3.45, 3.3, 3.45)
    C.text(s, 3.2, 2.55, 0.85, 0.55, ["HTTPS", "JWT"], size=10, color=PURPLE, bold=True, align="c")

    # 서버
    C.box(s, 4.0, 1.45, 3.0, 3.75, fill=NAVY)
    C.text(s, 4.2, 1.58, 2.7, 0.35, "Spring Boot 3.5", size=16, bold=True, color=WHITE)
    C.text(s, 4.2, 1.93, 2.7, 0.3, "Java 17 · Railway 운영 서버", size=10.5, color=LILAC2)
    mods = ["인증 · JWT 발급", "레슨 · 진도 · 발음(TTS)", "단어 · FSRS-4.5 스케줄러", "미션 대화 · 병렬 AI 호출", "실험실 · AI 응답 캐시"]
    y = 2.35
    for m in mods:
        C.box(s, 4.2, y, 2.6, 0.42, fill="40307F", radius=0.1)
        C.text(s, 4.35, y, 2.4, 0.42, m, size=11.5, color=WHITE, anchor="m")
        y += 0.5

    # DB
    C.arrow(s, 5.5, 5.22, 5.5, 5.55)
    db = C.box(s, 4.0, 5.55, 3.0, 1.35, fill=MINT, shape=MSO_SHAPE.CAN)
    db.adjustments[0] = 0.18
    C.text(s, 4.05, 5.82, 2.9, 0.35, "PostgreSQL", size=15, bold=True, color=NAVY, align="c")
    C.text(s, 4.05, 6.17, 2.9, 0.6, ["레슨 JSONB · 단어 5,561개", "AI 응답 캐시 · TTS 캐시"], size=10.5, color=NAVY, align="c")

    # 외부 서비스
    ext = [("OpenAI GPT-4o", "미션 대화 · 교정 · 판정"), ("OpenAI gpt-4o-mini-tts", "한국어 발음 음성"),
           ("Google Gemini 2.5 Flash", "AI 문법 실험실"), ("Google · Apple", "소셜 로그인 (토큰 검증)")]
    y = 1.45
    for t, d in ext:
        C.box(s, 7.65, y, 1.85, 0.8, fill=WHITE, line=PURPLE, line_w=1.2)
        C.text(s, 7.75, y + 0.08, 1.7, 0.35, t, size=10.5, bold=True, color=PURPLE)
        C.text(s, 7.75, y + 0.42, 1.7, 0.3, d, size=10, color=GRAY)
        C.arrow(s, 7.02, y + 0.4, 7.62, y + 0.4, color=GRAY, w=1.2)
        y += 0.95
    C.text(s, 7.65, 5.35, 1.85, 0.3, "외부 AI · 인증", size=10, color=GRAY, align="c")

    # 오프라인 파이프라인
    C.box(s, 0.5, 5.55, 3.25, 1.35, fill=WHITE, line=YELLOW, line_w=2)
    C.text(s, 0.68, 5.62, 3.0, 0.3, "오프라인 콘텐츠 파이프라인", size=12.5, bold=True, color=NAVY)
    C.text(s, 0.68, 5.95, 3.0, 0.9, [
        "레슨: CSV → Kiwi 형태소 분석 → 레슨 JSON",
        "단어: 국립국어원 어휘 → Gemini 번역·예문",
        [("→ 멱등 SQL 패치로 DB 반영", {"bold": True, "color": PURPLE})],
    ], size=10.5, color=INK, line=1.25)
    C.arrow(s, 3.78, 6.22, 3.98, 6.22)
    # 캐릭터 장식
    C.image(s, ch["turtle_idle"], 8.15, 5.7, h=1.2)


# ---------------------------------------------------------------- 3
def chip_row(C, s, x, y, items, h=0.42, size=14, gap=0.08):
    for t, w, hl in items:
        C.label(s, x, y, w, h, t, fill=(PURPLE if hl else WHITE), color=(WHITE if hl else NAVY),
                size=size, line=(None if hl else LILAC2), radius=0.08)
        x += w + gap
    return x


def phone(C, s, path, x, y, w, cap=None):
    pic = C.image(s, path, x, y, w=w)
    if cap:
        h = pic.height / C.unit
        C.text(s, x - 0.15, y + h + 0.06, w + 0.3, 0.3, cap, size=10, color=GRAY, align="c")
    return pic


def stat(C, s, x, y, w, h, big, small, fill=LILAC, color=PURPLE, big_size=24):
    C.box(s, x, y, w, h, fill=fill)
    C.text(s, x + 0.15, y + 0.06, w - 0.3, h * 0.5, big, size=big_size, bold=True, color=color, anchor="b")
    C.text(s, x + 0.15, y + h * 0.6, w - 0.3, h * 0.38, small, size=10, color=INK, line=1.15)


def page3(C, s, ph, ch):
    phone(C, s, ph["lesson_hint"], 0.55, 1.42, 1.6, "거북이 단계 · 힌트 말풍선")
    phone(C, s, ph["lesson_done"], 2.35, 1.42, 1.6, "레슨 완료 · 토끼 환호")

    X = 4.25
    C.box(s, X, 1.42, 5.25, 1.2, fill=LILAC)
    C.image(s, ch["rabbit_talking"], X + 0.12, 1.47, h=1.1)
    C.text(s, X + 1.15, 1.5, 4.0, 0.35, [[("토끼 단계", {"color": PURPLE}), (" — 의미 덩어리로 문장 조립", {})]], size=14, bold=True, color=NAVY)
    chip_row(C, s, X + 1.15, 1.98, [("저는", 0.75, 0), ("커피를", 0.95, 0), ("좋아해요", 1.2, 0)])

    C.arrow(s, X + 2.6, 2.66, X + 2.6, 2.86)
    C.box(s, X, 2.9, 5.25, 1.45, fill=LILAC)
    C.image(s, ch["turtle_thinking"], X + 0.1, 2.98, h=1.15)
    C.text(s, X + 1.15, 2.98, 4.0, 0.35, [[("거북이 단계", {"color": PURPLE}), (" — 레슨 목표 문법만 분해", {})]], size=14, bold=True, color=NAVY)
    chip_row(C, s, X + 1.15, 3.43, [("저는", 0.75, 0), ("커피", 0.75, 0), ("를", 0.5, 1), ("좋아해요", 1.2, 0)])
    C.text(s, X + 1.15, 3.92, 4.0, 0.35, "‘저는’·‘좋아해요’는 통째로, 이번 레슨 목표(을/를)인 ‘커피를’만 분리",
           size=10.5, color=GRAY)

    # 파이프라인
    C.text(s, X, 4.5, 5.3, 0.3, "Kiwi 기반 오프라인 콘텐츠 파이프라인", size=13, bold=True, color=NAVY)
    flow = [("CSV", "문장+목표 품사"), ("Kiwi", "형태소 분석"), ("생성", "조립 문제·탭 분석"), ("SQL 패치", "멱등·DB 반영")]
    x = X
    for j, (a, b) in enumerate(flow):
        C.box(s, x, 4.85, 1.12, 0.75, fill=WHITE, line=PURPLE, line_w=1.2)
        C.text(s, x, 4.9, 1.12, 0.3, a, size=12, bold=True, color=PURPLE, align="c")
        C.text(s, x, 5.2, 1.12, 0.35, b, size=9, color=GRAY, align="c")
        if j < 3:
            C.arrow(s, x + 1.14, 5.22, x + 1.36, 5.22)
        x += 1.38

    y = 5.85
    stat(C, s, 0.5, y, 2.15, 1.05, "12개", "한글 레슨 · 자모·음절 80항목")
    stat(C, s, 2.78, y, 2.15, 1.05, "3개", "문법 레슨 · 은/는, 이에요/예요, 을/를")
    stat(C, s, 5.06, y, 2.15, 1.05, "첫 시도", "점수 = 첫 시도 정답률 · 별 · 이어하기")
    stat(C, s, 7.34, y, 2.16, 1.05, "TTS", "서버 음성 생성 · DB 캐시 재사용")


# ---------------------------------------------------------------- 4
def page4(C, s, ph, ch):
    # 망각 곡선 차트
    cd = XyChartData()
    for S, name in ((1, "S = 1일"), (4, "S = 4일"), (15, "S = 15일")):
        ser = cd.add_series(name)
        for t in [i * 0.5 for i in range(0, 61)]:
            ser.add_data_point(t, (1 + 19 / 81 * t / S) ** -0.5)
    C.box(s, 0.5, 1.42, 3.75, 3.0, fill=WHITE, line=LILAC2, line_w=1.2)
    C.text(s, 0.68, 1.5, 3.5, 0.3, "기억률 곡선  R(t) = (1 + 19/81 · t/S)^−0.5", size=11, bold=True, color=NAVY)
    C.text(s, 0.68, 1.78, 3.5, 0.25, "t: 경과 일수 · S: 안정성 · R = 90%가 되는 날이 다음 복습일", size=9, color=GRAY)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS, C.e(0.6), C.e(2.02), C.e(3.55), C.e(2.35), cd)
    chart = gf.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9 * C.fs)
    chart.legend.font.name = FONT
    for ser, col in zip(chart.plots[0].series, (PINK, YELLOW, PURPLE)):
        ser.format.line.color.rgb = RGBColor.from_string(col)
        ser.format.line.width = Pt(2.2 * C.fs)
        ser.smooth = True
    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0.5, 1.0, 0.1
    ca.minimum_scale, ca.maximum_scale, ca.major_unit = 0, 30, 5
    for ax in (va, ca):
        ax.tick_labels.font.size = Pt(8.5 * C.fs)
        ax.tick_labels.font.color.rgb = RGBColor.from_string(GRAY)
        ax.format.line.color.rgb = RGBColor.from_string(LILAC2)
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string("ECEAF5")
    va.tick_labels.number_format = '0%'
    va.tick_labels.number_format_is_linked = False
    ca.has_major_gridlines = False

    # S D R
    defs = [("S", "안정성", "기억이 유지되는 기간. 잘 기억할수록 커진다."),
            ("D", "난이도", "단어별 어려움(1~10). 평가에 따라 조정."),
            ("R", "인출가능성", "지금 떠올릴 확률. 90% 직전이 복습일.")]
    y = 1.42
    for k, n, d in defs:
        C.box(s, 4.4, y, 2.3, 0.93, fill=LILAC)
        C.label(s, 4.52, y + 0.24, 0.45, 0.45, k, fill=PURPLE, size=15, radius=0.225)
        C.text(s, 5.08, y + 0.1, 1.55, 0.3, n, size=12.5, bold=True, color=NAVY)
        C.text(s, 5.08, y + 0.42, 1.55, 0.48, d, size=9.5, color=INK, line=1.15)
        y += 1.035

    # 평가 버튼
    C.text(s, 0.5, 4.6, 6.2, 0.35, "4단계 평가 → FSRS-4.5 표준 공식 · 기본 파라미터로 다음 복습일 계산", size=13, bold=True, color=NAVY)
    btn = [("Again", RED), ("Hard", YELLOW), ("Good", MINT), ("Easy", PURPLE)]
    x = 0.5
    for t, c in btn:
        C.label(s, x, 5.02, 1.45, 0.45, t, fill=c, size=13, radius=0.12)
        x += 1.55
    C.text(s, 0.5, 5.52, 6.2, 0.3, "버튼마다 다음 간격 미리보기 (예: Again 5분 · Good 4일) · ‘오늘의 복습’은 기한이 된 단어만",
           size=10, color=GRAY)

    y = 5.95
    stat(C, s, 0.5, y, 1.45, 0.95, "5,561", "학습 단어", big_size=20)
    stat(C, s, 2.05, y, 1.45, 0.95, "14개", "주제 덱", big_size=20)
    stat(C, s, 3.6, y, 1.45, 0.95, "30단어", "레슨 단위", big_size=20)
    stat(C, s, 5.15, y, 1.5, 0.95, "5쌍", "짝맞추기 라운드", big_size=20)

    phone(C, s, ph["vocab_review"], 6.95, 1.42, 1.18, "오늘의 복습 완료")
    phone(C, s, ph["vocab_match"], 8.3, 1.42, 1.18, "짝맞추기 완료")
    C.text(s, 6.9, 4.42, 2.6, 1.3, ["국립국어원 학습용 어휘 기반,", "번역·예문은 Gemini로 생성해", "주제별 덱으로 구성"],
           size=10.5, color=INK, line=1.3)
    C.image(s, ch["turtle_happy"], 7.55, 5.55, h=1.38)


# ---------------------------------------------------------------- 5
def page5(C, s, ph, ch):
    # 미션 흐름
    C.text(s, 0.5, 1.38, 5.9, 0.35, "Roleplay 미션 — 목적 · 끝 · 피드백이 있는 AI 대화", size=14, bold=True, color=NAVY)
    C.box(s, 0.5, 1.95, 1.25, 0.95, fill=LILAC)
    C.text(s, 0.5, 1.95, 1.25, 0.95, ["학습자", "한국어 문장"], size=11.5, bold=True, color=NAVY, align="c", anchor="m")
    C.arrow(s, 1.78, 2.2, 2.25, 1.98)
    C.arrow(s, 1.78, 2.65, 2.25, 2.87)
    C.box(s, 2.3, 1.75, 2.1, 0.55, fill=WHITE, line=PINK, line_w=1.5)
    C.image(s, ch["rabbit_talking"], 2.36, 1.78, h=0.5)
    C.text(s, 2.85, 1.75, 1.55, 0.55, "토끼 · 역할 연기", size=11.5, bold=True, color=NAVY, anchor="m")
    C.box(s, 2.3, 2.6, 2.1, 0.55, fill=WHITE, line=MINT, line_w=1.5)
    C.image(s, ch["turtle_talking"], 2.36, 2.63, h=0.5)
    C.text(s, 2.85, 2.6, 1.55, 0.55, "거북이 · 채점 교정", size=11.5, bold=True, color=NAVY, anchor="m")
    C.text(s, 2.3, 2.3, 2.1, 0.3, "동시에 호출 (GPT-4o)", size=9.5, color=PURPLE, bold=True, align="c", anchor="m")
    C.arrow(s, 4.43, 2.02, 4.85, 2.25)
    C.arrow(s, 4.43, 2.87, 4.85, 2.65)
    C.box(s, 4.9, 1.9, 1.5, 1.05, fill=PURPLE)
    C.text(s, 4.9, 1.9, 1.5, 1.05, [[("한 턴", {"size": 10.5})], [("1.5~2초", {"size": 18})], [("순차 대비 30~50%↓", {"size": 9.5})]],
           bold=True, color=WHITE, align="c", anchor="m", line=1.05)

    # 3-Zone
    C.text(s, 0.5, 3.2, 5.9, 0.3, "3-Zone 판정 — 서버가 강제, 최대 minTurns+3 턴 안에 종료", size=12, bold=True, color=NAVY)
    zones = [("A 초반", "판정 보류 · 대화 유지", LILAC2, NAVY), ("B 판정 구간", "목표 달성 여부 판단", PURPLE, WHITE), ("C 연장", "턴 상한 도달 → 종료", NAVY, WHITE)]
    x = 0.5
    for j, (a, b, f, c) in enumerate(zones):
        C.box(s, x, 3.55, 1.93, 0.62, fill=f, radius=0.08)
        C.text(s, x, 3.57, 1.93, 0.3, a, size=11.5, bold=True, color=c, align="c")
        C.text(s, x, 3.85, 1.93, 0.3, b, size=9.5, color=c, align="c")
        x += 1.985

    # 폰 3장 + 설명
    phone(C, s, ph["mission_magic"], 0.5, 4.35, 1.08, "역할로 변신 로딩")
    phone(C, s, ph["mission_red"], 1.72, 4.35, 1.08, "레드카드 교정")
    phone(C, s, ph["mission_report"], 2.94, 4.35, 1.08, "미션 리포트")
    C.text(s, 4.2, 4.4, 2.2, 2.5, [
        [("레드 · 옐로카드", {"bold": True, "color": PURPLE})],
        "대화 중 고칠 표현을 카드로 바로 알려 준다.",
        [("AI 코치 판정", {"bold": True, "color": PURPLE, "space_before": 6})],
        "목표 달성 시 클리어, 미달 시 재도전 권유.",
        [("항상 피드백", {"bold": True, "color": PURPLE, "space_before": 6})],
        "잘한 표현 · 고칠 표현 · 다음 목표 (실제로 보낸 문장만 인용)",
    ], size=10.5, color=INK, line=1.2)

    # 실험실
    X = 6.65
    C.box(s, X, 1.38, 2.85, 5.52, fill=LILAC)
    C.text(s, X + 0.18, 1.5, 2.5, 0.35, "Upgrade 실험실", size=14, bold=True, color=NAVY)
    C.text(s, X + 0.18, 1.9, 2.5, 0.3, "한글랩 — 자모 → 글자 조합", size=11, bold=True, color=PURPLE)
    C.box(s, X + 0.18, 2.22, 2.5, 0.5, fill=WHITE, radius=0.08)
    C.text(s, X + 0.18, 2.22, 2.5, 0.5, "0xAC00 + (초성×21 + 중성)×28 + 종성", size=9.5, color=NAVY, bold=True, align="c", anchor="m")
    C.text(s, X + 0.18, 2.75, 2.5, 0.3, "유니코드 공식 · 기기에서 즉시 계산", size=9.5, color=GRAY)
    C.text(s, X + 0.18, 3.12, 2.5, 0.3, "그래머랩 — AI 문법 변형", size=11, bold=True, color=PURPLE)
    C.text(s, X + 0.18, 3.42, 2.5, 0.3, "Gemini 2.5 Flash · 같은 요청은 DB 캐시", size=9.5, color=GRAY)
    C.box(s, X + 0.18, 3.78, 1.2, 0.85, fill=WHITE)
    C.text(s, X + 0.18, 3.8, 1.2, 0.5, "1~2ms", size=17, bold=True, color=PURPLE, align="c", anchor="m")
    C.text(s, X + 0.18, 4.25, 1.2, 0.3, "캐시 HIT", size=9.5, color=GRAY, align="c")
    C.box(s, X + 1.48, 3.78, 1.2, 0.85, fill=WHITE)
    C.text(s, X + 1.48, 3.8, 1.2, 0.5, "3~5초", size=17, bold=True, color=NAVY, align="c", anchor="m")
    C.text(s, X + 1.48, 4.25, 1.2, 0.3, "MISS (AI 호출)", size=9.5, color=GRAY, align="c")
    C.text(s, X + 0.18, 4.7, 2.5, 0.3, "수식어 순서가 달라도 같은 캐시 키", size=9.5, color=GRAY)
    phone(C, s, ph["lab_grammar"], X + 0.18, 5.05, 0.82)
    C.image(s, ch["turtle_happy"], X + 1.3, 5.45, h=1.35)


# ---------------------------------------------------------------- 6
def page6(C, s, ph, ch):
    C.box(s, 0.5, 1.42, 3.7, 3.25, fill=LILAC)
    C.image(s, ch["rabbit_idle"], 0.6, 1.55, w=1.7, h=2.45)
    C.image(s, ch["turtle_idle"], 2.35, 1.75, w=1.75, h=2.25)
    C.text(s, 0.6, 4.05, 1.7, 0.55, [[("토끼", {"bold": True, "color": PURPLE})], "빠른 체험 · 연기"], size=11, color=INK, align="c", line=1.1)
    C.text(s, 2.35, 4.05, 1.75, 0.55, [[("거북이", {"bold": True, "color": PURPLE})], "느린 분석 · 코칭"], size=11, color=INK, align="c", line=1.1)

    C.text(s, 4.45, 1.38, 5.0, 0.35, "학습 이벤트에 실시간 반응 — 4개 기능 21개 화면 지점 + 홈", size=13, bold=True, color=NAVY)
    rows = [("rabbit_happy", "정답", "캐릭터가 함께 기뻐한다"),
            ("turtle_thinking", "오답", "거북이가 힌트 말풍선으로 코칭"),
            ("rabbit_cheer", "레슨 완료", "토끼 환호 점프 + 축하 효과"),
            ("rabbit_magic", "미션 준비", "토끼가 역할로 변신하는 로딩"),
            ("turtle_talking", "대화 교정", "거북이가 레드 · 옐로카드로 교정"),
            ("rabbit_blink", "대기", "숨쉬기 · 눈 깜빡임")]
    y = 1.8
    for j, (img, ev, re) in enumerate(rows):
        x = 4.45 if j % 2 == 0 else 7.0
        if j % 2 == 0 and j:
            y += 0.97
        C.box(s, x, y, 2.45, 0.85, fill=WHITE, line=LILAC2, line_w=1.2)
        C.image(s, ch[img], x + 0.08, y + 0.06, w=0.7, h=0.73)
        C.text(s, x + 0.85, y + 0.08, 1.55, 0.3, ev, size=11.5, bold=True, color=PURPLE)
        C.text(s, x + 0.85, y + 0.38, 1.55, 0.45, re, size=9.5, color=INK, line=1.1)

    # 표정 스트립
    order = ["rabbit_idle", "rabbit_happy", "rabbit_cheer", "rabbit_thinking", "rabbit_talking", "rabbit_sad", "rabbit_blink", "rabbit_magic",
             "turtle_idle", "turtle_happy", "turtle_cheer", "turtle_thinking", "turtle_talking", "turtle_sad", "turtle_blink"]
    C.text(s, 0.5, 4.82, 9.0, 0.3, "AI 생성 표정 이미지 15장 (토끼 8 · 거북이 7)", size=11, bold=True, color=NAVY)
    w = 9.0 / 15
    for j, k in enumerate(order):
        C.image(s, ch[k], 0.5 + j * w + 0.02, 5.15, w=w - 0.04, h=0.62)

    # 설계
    y = 5.98
    C.box(s, 0.5, y, 2.1, 0.9, fill=LILAC)
    C.text(s, 0.5, y, 2.1, 0.9, [[("표정 그림", {"bold": True, "color": PURPLE})], "AI 생성 PNG"], size=11.5, color=INK, align="c", anchor="m", line=1.15)
    C.text(s, 2.6, y, 0.4, 0.9, "+", size=20, bold=True, color=PURPLE, align="c", anchor="m")
    C.box(s, 3.0, y, 2.4, 0.9, fill=LILAC)
    C.text(s, 3.0, y, 2.4, 0.9, [[("몸짓", {"bold": True, "color": PURPLE})], "코드 모션: 숨쉬기·깜빡임·점프·찌그러짐"], size=9.5, color=INK, align="c", anchor="m", line=1.15)
    C.arrow(s, 5.45, y + 0.45, 5.85, y + 0.45)
    C.box(s, 5.9, y, 3.6, 0.9, fill=PURPLE)
    C.text(s, 5.9, y, 3.6, 0.9, [[("그림 교체 시 기능 코드 수정 0줄", {"bold": True})], "없는 표정은 자동 폴백"], size=12, color=WHITE, align="c", anchor="m", line=1.15)


# ---------------------------------------------------------------- 7
def page7(C, s, ch):
    C.text(s, 0.5, 1.35, 9.0, 0.35, "Railway 운영 서버 배포 · iPhone 실기기 설치와 Google 로그인까지 확인한 실제 서비스",
           size=13, color=NAVY, bold=True)
    sts = [("4개", "핵심 학습 기능 (M · A · R · U)"), ("15개", "레슨 (한글 12 + 문법 3)"),
           ("5,561", "학습 단어 · 14개 주제 덱"), ("123 · 67", "자동 테스트 (서버 · 앱) 통과")]
    x = 0.5
    for b, sm in sts:
        stat(C, s, x, 1.85, 2.15, 1.1, b, sm, big_size=24)
        x += 2.283

    cols = [("학습자", YELLOW, "rabbit_happy", [
                "덩어리 → 목표 문법 분해로 교착어 구조를 직접 체험",
                "잊기 직전에 복습하는 FSRS-4.5로 효율적인 단어 기억",
                "목적 · 종료 · 피드백이 있는 AI 회화 연습"]),
            ("교육 · 콘텐츠", MINT, "turtle_thinking", [
                "CSV(문장 + 목표 품사)만으로 레슨 데이터 생성",
                "어휘 파이프라인 재사용으로 주제 덱 확장",
                "멱등 SQL 패치로 운영 DB에 안전하게 반영"]),
            ("개발 · 운영", PINK, "turtle_happy", [
                "AI 응답 캐시 · 병렬 호출로 응답 시간과 비용 관리",
                "캐릭터 그림 / 모션 분리로 디자인 교체 비용 최소화",
                "자동 테스트 + 푸시 시 Railway 자동 배포"])]
    x = 0.5
    for t, c, img, bl in cols:
        C.box(s, x, 3.15, 2.9, 2.6, fill=WHITE, line=LILAC2, line_w=1.5)
        C.box(s, x + 0.2, 3.32, 0.62, 0.62, fill=c, shape=MSO_SHAPE.OVAL)
        C.image(s, ch[img], x + 0.22, 3.33, w=0.58, h=0.6)
        C.text(s, x + 0.95, 3.32, 1.9, 0.62, t, size=15, bold=True, color=NAVY, anchor="m")
        paras = []
        for k, b in enumerate(bl):
            paras.append([("· " + b, {"space_before": 5 if k else 0})])
        C.text(s, x + 0.22, 4.1, 2.5, 1.6, paras, size=11, color=INK, line=1.2)
        x += 3.05

    C.box(s, 0.5, 5.95, 9.0, 0.95, fill=PURPLE)
    C.image(s, ch["rabbit_cheer"], 0.65, 5.85, h=1.05)
    C.image(s, ch["turtle_cheer"], 8.2, 5.88, h=1.02)
    C.text(s, 1.9, 5.95, 6.2, 0.95, [[("향후 계획  ", {"color": YELLOW}), ("웹 체험판 · 랜딩 페이지 → 2026년 11월 졸업전시", {})]],
           size=15, bold=True, color=WHITE, align="c", anchor="m")


if __name__ == "__main__":
    main()
