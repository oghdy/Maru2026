"""dlv-design 공통 도우미 — 양식 pptx 위에 python-pptx 로 도형·글·이미지를 올린다.

좌표는 '가상 그리드' 단위(gx, gy)로 받고, 실제 슬라이드 크기에 맞춰 SCALE 배로 키운다.
글꼴 크기도 같은 SCALE 로 키워 '보통 크기 슬라이드에서 설계한 대로' 보이게 한다.
"""
import copy
import os
from PIL import Image, ImageDraw, ImageFilter
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

FONT = "맑은 고딕"

# 브랜드 팔레트 (PLAN.md 브랜드)
PURPLE = "6B4EFF"
NAVY = "2B1D5C"
INK = "2B2540"
GRAY = "6E6889"
LILAC = "F0EDFF"   # 연보라 카드
LILAC2 = "E2DCFF"
YELLOW = "FFC83D"
MINT = "4CD4B0"
PINK = "FF8FB1"
RED = "F0506E"
WHITE = "FFFFFF"


class Canvas:
    def __init__(self, prs, grid_w):
        self.prs = prs
        self.unit = prs.slide_width / grid_w           # EMU per grid unit
        self.fs = self.unit / 914400                   # 글꼴 배율 (1 grid unit = 1 inch 기준)

    def e(self, v):
        return Emu(int(round(v * self.unit)))

    # ---------- 글 ----------
    def text(self, slide, x, y, w, h, paras, size=14, color=INK, bold=False,
             align="l", anchor="t", line=1.15, margin=0, name=None):
        """paras: str | list[str | list[run]], run = (text, {size,color,bold})"""
        tb = slide.shapes.add_textbox(self.e(x), self.e(y), self.e(w), self.e(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        m = self.e(margin)
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, str):
            paras = [paras]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            para.line_spacing = line
            runs = p if isinstance(p, list) else [(p, {})]
            for rt, ro in runs:
                r = para.add_run()
                r.text = rt
                self.style_run(r, ro.get("size", size), ro.get("color", color), ro.get("bold", bold))
            if isinstance(p, list) and p and p[0][1].get("space_before"):
                para.space_before = Pt(p[0][1]["space_before"] * self.fs)
        return tb

    def style_run(self, r, size, color, bold):
        f = r.font
        f.size = Pt(size * self.fs)
        f.bold = bold
        f.color.rgb = RGBColor.from_string(color)
        f.name = FONT
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", FONT)

    # ---------- 도형 ----------
    def box(self, slide, x, y, w, h, fill=LILAC, line=None, radius=0.12, shape=None, line_w=1.0, shadow=False):
        st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        s = slide.shapes.add_shape(st, self.e(x), self.e(y), self.e(w), self.e(h))
        if st == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = min(0.5, radius / min(w, h))
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(line_w * self.fs)
        if not shadow:
            sp = s._element.spPr
            el = etree.SubElement(sp, qn("a:effectLst"))
        s.text_frame.text = ""
        return s

    def label(self, slide, x, y, w, h, txt, fill=PURPLE, color=WHITE, size=12, bold=True, radius=None, align="c", line=None):
        s = self.box(slide, x, y, w, h, fill=fill, radius=(h / 2 if radius is None else radius), line=line)
        tf = s.text_frame
        tf.word_wrap = True
        z = self.e(0.04)
        tf.margin_left = tf.margin_right = z
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = txt if isinstance(txt, list) else [txt]
        for i, t in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            runs = t if isinstance(t, list) else [(t, {})]
            for rt, ro in runs:
                r = p.add_run()
                r.text = rt
                self.style_run(r, ro.get("size", size), ro.get("color", color), ro.get("bold", bold))
        return s

    def arrow(self, slide, x1, y1, x2, y2, color=PURPLE, w=2.0, head=True):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.e(x1), self.e(y1), self.e(x2), self.e(y2))
        c.line.color.rgb = RGBColor.from_string(color)
        c.line.width = Pt(w * self.fs)
        if head:
            ln = c.line._get_or_add_ln()
            te = etree.SubElement(ln, qn("a:tailEnd"))
            te.set("type", "triangle")
            te.set("w", "med")
            te.set("h", "med")
        return c

    def image(self, slide, path, x, y, w=None, h=None, name=None):
        im = Image.open(path)
        iw, ih = im.size
        if w is not None and h is not None:  # 맞춤(contain), 가운데 정렬
            r = min(w / iw, h / ih)
            nw, nh = iw * r, ih * r
            x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
        elif w is not None:
            h = w * ih / iw
        else:
            w = h * iw / ih
        pic = slide.shapes.add_picture(path, self.e(x), self.e(y), self.e(w), self.e(h))
        if name:
            pic.name = name
        return pic


# ---------- 이미지 가공 ----------
def trim_alpha(src, dst, max_side=None):
    im = Image.open(src).convert("RGBA")
    im = im.crop(im.getchannel("A").getbbox())
    if max_side and max(im.size) > max_side:
        im.thumbnail((max_side, max_side), Image.LANCZOS)
    im.save(dst, optimize=True)
    return dst


def phone_frame(src, dst, width=720, bezel=0.03, crop_top=0.0):
    """스크린샷 → 둥근 모서리 + 남색 테두리 폰 프레임 PNG(투명 배경)."""
    sc = Image.open(src).convert("RGB")
    if crop_top:
        sc = sc.crop((0, int(sc.height * crop_top), sc.width, sc.height))
    inner_w = int(width * (1 - 2 * bezel))
    sc = sc.resize((inner_w, int(sc.height * inner_w / sc.width)), Image.LANCZOS)
    b = (width - inner_w) // 2
    W, H = width, sc.height + 2 * b
    out = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    r_out = int(W * 0.13)
    m = Image.new("L", (W, H), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, W - 1, H - 1), r_out, fill=255)
    frame = Image.new("RGBA", (W, H), (0x2B, 0x1D, 0x5C, 255))
    out.paste(frame, (0, 0), m)
    mi = Image.new("L", sc.size, 0)
    ImageDraw.Draw(mi).rounded_rectangle((0, 0, sc.width - 1, sc.height - 1), r_out - b, fill=255)
    out.paste(sc, (b, b), mi)
    out.save(dst, optimize=True)
    return dst


def placeholder_phone(dst, label, width=720, ratio=2622 / 1206):
    H = int(width * ratio)
    out = Image.new("RGBA", (width, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(out)
    d.rounded_rectangle((0, 0, width - 1, H - 1), int(width * 0.13), fill=(240, 237, 255, 255),
                        outline=(107, 78, 255, 255), width=8)
    out.save(dst)
    return dst


# ---------- 슬라이드 조작 ----------
def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)


def move_slide(prs, old, new):
    sldIdLst = prs.slides._sldIdLst
    el = sldIdLst[old]
    sldIdLst.remove(el)
    sldIdLst.insert(new, el)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)
