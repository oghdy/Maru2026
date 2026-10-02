"""포스터 생성 — 양식(첨부2, 42×60cm) 1장, 시안 B '보라 히어로'. 실행: python3 build_poster.py
좌표 단위 = cm, 글꼴 = 실제 pt. 스크린샷 교체: SHOTS 경로만 바꾸고 다시 실행.
QR: 이름이 'QR_자리' 인 흰 사각형 위치(우상단)에 QR 이미지를 올리면 된다.
"""
import os
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "../../book/_build"))
from pptlib import *  # noqa

ROOT = os.path.abspath(os.path.join(HERE, "../../../.."))
DOCS = os.path.join(ROOT, "docs")
TEMPLATE = os.path.join(DOCS, "deliverables/templates/첨부2_포스터양식.pptx")
OUT = os.path.join(DOCS, "deliverables/poster/MARU_포스터.pptx")
ASSETS = os.path.join(HERE, "assets")
os.makedirs(ASSETS, exist_ok=True)
CH = os.path.join(DOCS, "features/character/screenshots")
MS = os.path.join(DOCS, "features/mission/screenshots")
RAW = os.path.join(DOCS, "features/character/raw")

SHOTS = {
    "lesson_hint": f"{CH}/apply_lesson_C1.png",
    "lesson_done": f"{CH}/apply_lesson_C3.png",
    "vocab_review": f"{CH}/apply_vocab_C3_review_done.png",
    "vocab_match": f"{CH}/apply_vocab_C2.png",
    "mission_magic": f"{MS}/r3_fe2/12_loading_magic.png",
    "mission_red": f"{MS}/MSN-1.7.3_red_card.png",
    "lab_hangeul": f"{CH}/apply_lab_C1.png",
    "lab_grammar": f"{CH}/apply_lab_C3.png",
}
DEEP = "5A3EE6"


def main():
    ph = {k: phone_frame(p, f"{ASSETS}/ph_{k}.png", width=900) for k, p in SHOTS.items()}
    ch = {f[:-4]: trim_alpha(f"{RAW}/{f}", f"{ASSETS}/c_{f}") for f in os.listdir(RAW) if f.endswith(".png")}

    prs = Presentation(TEMPLATE)
    C = Canvas(prs, 42.0)   # 1 unit = 1cm
    C.fs = 1.0              # 글꼴은 실제 pt
    # 안내(2번) 장 삭제, 1번 장의 예시 그림 제거
    delete_slide(prs, 1)
    s = prs.slides[0]
    for shp in list(s.shapes):
        remove_shape(shp)

    # ---------- 히어로 ----------
    C.box(s, 0, 0, 42, 25.5, fill=PURPLE, radius=0)
    C.box(s, 24, -9, 26, 26, fill="7A60FF", shape=MSO_SHAPE.OVAL)       # 배경 원 장식
    C.box(s, -6, 16, 14, 14, fill=DEEP, shape=MSO_SHAPE.OVAL)
    C.box(s, 0, 25.5, 42, 34.5, fill=WHITE, radius=0)                   # 아래 흰 바탕 (원 장식 가림)

    C.label(s, 2.5, 2.4, 8.2, 1.5, "TEAM  MARU", fill=WHITE, color=PURPLE, size=30)
    C.text(s, 11.4, 2.4, 20, 1.5, [[("팀원  ", {"color": LILAC2}), ("하도윤", {"bold": True}), ("  컴퓨터공학과", {"color": LILAC2, "size": 24})]],
           size=32, color=WHITE, anchor="m")
    C.text(s, 2.5, 4.6, 20, 1.0, "2026년 졸업작품 · 지도교수 한경수 교수", size=24, color=LILAC2)

    qr = C.box(s, 33.3, 2.0, 6.7, 6.7, fill=WHITE, radius=0.5)
    qr.name = "QR_자리"
    C.text(s, 33.3, 8.9, 6.7, 0.8, "팀 페이지", size=20, color=LILAC2, align="c")

    C.text(s, 2.0, 6.4, 30, 7.5, "MARU", size=220, bold=True, color=WHITE, line=0.9)
    C.text(s, 2.6, 13.9, 20.5, 2.0, [[("마루", {"bold": True, "color": YELLOW}), (" · 외국인을 위한 한국어 학습 앱", {})]],
           size=32, color=WHITE)
    C.text(s, 2.6, 17.0, 18.5, 4.5, [
        [("토끼처럼 ", {}), ("빠르게 체험", {"color": YELLOW, "bold": True}), ("하고,", {})],
        [("거북이처럼 ", {}), ("천천히 이해", {"color": MINT, "bold": True}), ("한다.", {})],
    ], size=40, color=WHITE, line=1.2)

    C.image(s, ch["rabbit_cheer"], 20.0, 11.2, h=15.6)
    C.image(s, ch["turtle_happy"], 30.2, 14.0, h=13.0)

    # ---------- 소개 ----------
    C.label(s, 2.5, 28.0, 10.0, 1.5, "INTRODUCTION", fill=PURPLE, size=30)
    C.text(s, 2.5, 30.0, 37, 4.0, [
        [("한국어는 단어에 조사·어미를 붙여 뜻을 만드는 ", {}), ("교착어", {"bold": True, "color": PURPLE}),
         ("다. 기존 학습 앱은 암기와 객관식 위주라 이 구조를 체험하기 어렵다. ", {}),
         ("MARU", {"bold": True, "color": PURPLE}),
         ("는 토끼와 함께 문장을 덩어리로 먼저 써 보고, 거북이와 함께 레슨 목표 문법만 쪼개 보며 한국어의 구조를 직접 익히는 앱이다.", {})],
    ], size=27, color=INK, line=1.35)

    # ---------- 4기능 2×2 ----------
    cards = [
        ("M", "Morpheme 레슨", ["덩어리로 조립(토끼) →", "목표 문법만 분해(거북이)", "Kiwi 콘텐츠 파이프라인", "한글 12 + 문법 3 레슨"],
         ["lesson_hint", "lesson_done"], YELLOW),
        ("A", "Adaptive 단어", ["FSRS-4.5 표준 공식으로", "단어별 다음 복습일 계산", "5,561단어 · 14개 주제 덱", "짝맞추기 라운드당 5쌍"],
         ["vocab_review", "vocab_match"], MINT),
        ("R", "Roleplay 미션", ["토끼가 역할로 변신해 대화", "거북이가 레드·옐로카드 교정", "한 턴 약 1.5~2초 (병렬 호출)", "AI 코치가 클리어 판정"],
         ["mission_magic", "mission_red"], PINK),
        ("U", "Upgrade 실험실", ["한글 자모 → 글자 조합", "AI 문법 변형 (Gemini)", "같은 요청은 DB 캐시", "캐시 HIT 1~2ms"],
         ["lab_hangeul", "lab_grammar"], PURPLE),
    ]
    W, H = 18.0, 10.4
    for i, (k, t, lines, shots, col) in enumerate(cards):
        x = 2.5 + (i % 2) * (W + 1.0)
        y = 34.6 + (i // 2) * (H + 0.9)
        C.box(s, x, y, W, H, fill=LILAC, radius=0.8)
        C.label(s, x + 0.8, y + 0.8, 2.0, 2.0, k, fill=col, color=(NAVY if col in (YELLOW, MINT) else WHITE), size=48, radius=1.0)
        C.text(s, x + 0.8, y + 3.1, 9.2, 1.4, t, size=30, bold=True, color=NAVY)
        paras = []
        for j, ln in enumerate(lines):
            st = {"color": INK} if j < 2 else {"color": PURPLE, "bold": True, "space_before": 8 if j == 2 else 0}
            paras.append([(ln, st)])
        C.text(s, x + 0.8, y + 4.7, 9.3, 5.4, paras, size=19, line=1.25)
        px = x + W - 0.6
        for sh in reversed(shots):
            pw = 3.4
            px -= pw
            C.image(s, ph[sh], px, y + (H - pw * 2.2) / 2, w=pw)
            px -= 0.35

    # ---------- 기술 ----------
    C.box(s, 2.5, 56.4, 37, 2.6, fill=NAVY, radius=1.3)
    C.text(s, 3.4, 56.4, 5, 2.6, "TECH", size=28, bold=True, color=YELLOW, anchor="m")
    tech = ["Flutter", "Spring Boot", "PostgreSQL", "Railway", "FSRS-4.5", "Kiwi", "OpenAI GPT-4o", "Gemini 2.5"]
    x = 7.6
    for t in tech:
        w = len(t) * 0.31 + 0.8
        C.label(s, x, 57.05, w, 1.3, t, fill="40307F", color=WHITE, size=16, bold=False)
        x += w + 0.3

    prs.core_properties.title = "MARU (마루) — 졸업작품 포스터"
    prs.core_properties.author = "하도윤"
    prs.part._element.set("embedTrueTypeFonts", "1")  # PowerPoint 재저장 시 글꼴 포함
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
